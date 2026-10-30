"""Small openpyxl helpers so every workbook the tool writes looks the same."""

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

# Syz Visual Guidelines palette (see syz-formating/syz/tokens.py).
SPACE_BLUE = "202945"
TIGER_ORANGE = "FF6C0E"
PANEL_FILL = "F3F5FA"
BODY_INK = "2A3252"
MUTED_INK = "66708A"
HAIRLINE = "D2D7DF"
WARN_FILL = "FFF1E0"
ERROR_FILL = "FBE3E6"
FONT = "Nunito"

FMT = {
    "date": "DD.MM.YYYY",
    "month": "MMM YYYY",
    "amount": "#,##0",
    "amount2": "#,##0.00",
    "pct": "0.00%",
    "rate": "0.0000",
    "int": "0",
    "text": "@",
}

HEADER_FONT = Font(name=FONT, bold=True, color="FFFFFF")
BODY_FONT = Font(name=FONT, color=BODY_INK)
TITLE_FONT = Font(name=FONT, bold=True, size=16, color=SPACE_BLUE)
NOTE_FONT = Font(name=FONT, italic=True, size=9, color=MUTED_INK)
HEADER_FILL = PatternFill("solid", fgColor=SPACE_BLUE)
TOTAL_FILL = PatternFill("solid", fgColor=PANEL_FILL)
THIN = Side(style="thin", color=HAIRLINE)


def write_table(ws, df, start_row=1, start_col=1, formats=None, name=None,
                total_row=False, widths=None):
    """Write `df` with a styled header. Returns the last row written.

    formats: {column name: key of FMT}. name: Excel table name (adds filter
    buttons and banding; leave None for small summary blocks).
    """
    formats = formats or {}
    cols = list(df.columns)
    for j, col in enumerate(cols):
        c = ws.cell(start_row, start_col + j, col)
        c.font, c.fill = HEADER_FONT, HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    r = start_row
    for r, row in enumerate(df.itertuples(index=False), start_row + 1):
        for j, val in enumerate(row):
            if val is not None and val != val:      # NaN -> empty cell
                val = None
            c = ws.cell(r, start_col + j, val)
            c.font = BODY_FONT
            c.border = Border(bottom=THIN)
            f = formats.get(cols[j])
            if f:
                c.number_format = FMT[f]
    if total_row and len(df):
        r += 1
        for j, col in enumerate(cols):
            c = ws.cell(r, start_col + j)
            c.fill = TOTAL_FILL
            c.font = Font(name=FONT, bold=True, color=BODY_INK)
            if j == 0:
                c.value = "Total"
            elif formats.get(col) in ("amount", "amount2", "int"):
                L = get_column_letter(start_col + j)
                c.value = f"=SUBTOTAL(109,{L}{start_row + 1}:{L}{r - 1})"
                c.number_format = FMT[formats[col]]
    if name and len(df):
        ref = (f"{get_column_letter(start_col)}{start_row}:"
               f"{get_column_letter(start_col + len(cols) - 1)}{start_row + len(df)}")
        t = Table(displayName=name, ref=ref)
        t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
        ws.add_table(t)
    for j, col in enumerate(cols):
        L = get_column_letter(start_col + j)
        if widths and col in widths:
            ws.column_dimensions[L].width = widths[col]
        else:
            longest = max([len(str(col))] + [len(str(v)) for v in df[col].head(500)])
            ws.column_dimensions[L].width = min(max(longest + 2, 10), 45)
    return r


def title(ws, text, subtitle=None):
    ws["A1"] = text
    ws["A1"].font = TITLE_FONT
    if subtitle:
        ws["A2"] = subtitle
        ws["A2"].font = NOTE_FONT
    ws.sheet_view.showGridLines = False
