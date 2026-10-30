"""Build reoffer_report.xlsx from the register: CHF conversion and every
breakdown the review tables need. Rebuilt in full on each run, so it never
drifts from the register."""

import datetime as dt

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference as ChartRef

from . import excel
from .schema import HEADERS

MONEY = {"Trades": "int", "Notional (CHF)": "amount", "Fees (CHF)": "amount",
         "Avg fee": "pct", "Share of fees": "pct"}


def with_chf(register, ref) -> pd.DataFrame:
    df = register.copy()
    df["fx_to_chf"] = [ref.chf_rate(c, d) for c, d in zip(df["currency"], df["trade_date"])]
    df["size_chf"] = df["size"].astype(float) * df["fx_to_chf"]
    df["fee_chf"] = df["fee_amount"].astype(float) * df["fx_to_chf"]
    df["month"] = pd.to_datetime(df["trade_date"]).dt.to_period("M").dt.to_timestamp()
    df["year"] = pd.to_datetime(df["trade_date"]).dt.year
    return df


def breakdown(df, by, label):
    g = df.groupby(by, dropna=False).agg(
        trades=("trade_id", "count"), notional=("size_chf", "sum"),
        fees=("fee_chf", "sum")).reset_index()
    g["avg"] = g["fees"] / g["notional"]
    total = g["fees"].sum()
    g["share"] = g["fees"] / total if total else 0
    g = g.rename(columns={by: label, "trades": "Trades", "notional": "Notional (CHF)",
                          "fees": "Fees (CHF)", "avg": "Avg fee", "share": "Share of fees"})
    if label != "Month":
        g = g.sort_values("Fees (CHF)", ascending=False)
    return g


def build_report(register, ref, path):
    df = with_chf(register, ref)
    missing = df[df["fx_to_chf"].isna()]
    ok = df[df["fx_to_chf"].notna()]

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    excel.title(ws, "Reoffer fees — summary",
                f"Built {dt.datetime.now():%d.%m.%Y %H:%M} from {len(df)} register rows. "
                "All amounts in CHF.")
    kpis = pd.DataFrame([
        ("Trades", len(ok)),
        ("Notional (CHF)", ok["size_chf"].sum()),
        ("Fees (CHF)", ok["fee_chf"].sum()),
        ("Avg fee", ok["fee_chf"].sum() / ok["size_chf"].sum() if len(ok) else 0),
        ("Rows without FX rate (excluded)", len(missing)),
    ], columns=["Measure", "Value"])
    end = excel.write_table(ws, kpis, start_row=4, widths={"Measure": 34, "Value": 18})
    for r, fmt in zip(range(5, 10), ["int", "amount", "amount", "pct", "int"]):
        ws.cell(r, 2).number_format = excel.FMT[fmt]
    if len(missing):
        ws.cell(9, 2).fill = excel.PatternFill("solid", fgColor=excel.ERROR_FILL)
    ws.cell(end + 2, 1, "By year").font = excel.TITLE_FONT
    excel.write_table(ws, breakdown(ok, "year", "Year"), start_row=end + 3,
                      formats=MONEY, total_row=True)
    ws.column_dimensions["A"].width = 34

    sheets = [("By Month", "month", "Month"), ("By Issuer", "issuer", "Issuer"),
              ("By RM", "rm", "RM"), ("By Region", "region", "Region"),
              ("By Currency", "currency", "Currency"),
              ("By Mgmt Style", "management_style", "Management Style"),
              ("By Pay-off", "payoff", "Pay-off"), ("By APL", "apl", "APL")]
    for name, by, label in sheets:
        w = wb.create_sheet(name)
        excel.title(w, name.replace("By ", "Fees by "))
        t = breakdown(ok, by, label)
        last = excel.write_table(w, t, start_row=4, formats={**MONEY, "Month": "month"},
                                 total_row=True)
        w.freeze_panes = "B5"
        if len(t) > 1:
            _bar(w, t, label, rows=(5, 4 + len(t)), anchor=f"I4")

    # RM x Month matrix (fees CHF): the monthly RM review in one grid.
    w = wb.create_sheet("RM x Month")
    excel.title(w, "Fees (CHF) by RM and month")
    piv = ok.pivot_table(index="rm", columns="month", values="fee_chf",
                         aggfunc="sum", fill_value=0)
    piv.columns = [c.strftime("%b %Y") for c in piv.columns]
    piv["Total"] = piv.sum(axis=1)
    piv = piv.sort_values("Total", ascending=False).reset_index().rename(columns={"rm": "RM"})
    excel.write_table(w, piv, start_row=4, total_row=True,
                      formats={c: "amount" for c in piv.columns if c != "RM"})
    w.freeze_panes = "B5"

    # Full register with CHF columns, for ad-hoc pivots.
    w = wb.create_sheet("Data")
    cols = ["trade_id"] + list(HEADERS) + ["fx_to_chf", "size_chf", "fee_chf", "month"]
    out = df[cols].rename(columns={**HEADERS, "trade_id": "Trade ID", "fx_to_chf": "FX to CHF",
                                   "size_chf": "Size (CHF)", "fee_chf": "Fees (CHF)",
                                   "month": "Month"})
    excel.write_table(w, out, name="Data", formats={
        "Trade date": "date", "Size": "amount", "Fees (%)": "pct",
        "Fees (Amount)": "amount2", "FX to CHF": "rate", "Size (CHF)": "amount",
        "Fees (CHF)": "amount", "Month": "month"}, widths={"Comments": 40})
    w.freeze_panes = "C2"

    if len(missing):
        w = wb.create_sheet("Missing FX")
        excel.title(w, "Rows excluded: no FX rate",
                    "Add the rate to reference/fx_rates.csv (CHF per 1 unit, "
                    "dated on or before the trade date) and re-run.")
        m = (missing.groupby("currency")["trade_date"].agg(["count", "min", "max"])
             .reset_index().rename(columns={"currency": "Currency", "count": "Rows",
                                            "min": "First trade", "max": "Last trade"}))
        excel.write_table(w, m, start_row=4, formats={"First trade": "date",
                                                      "Last trade": "date"})
    wb.save(path)
    return path


def _bar(ws, t, label, rows, anchor):
    fees_col = list(t.columns).index("Fees (CHF)") + 1
    chart = BarChart()
    chart.type = "col" if label == "Month" else "bar"
    chart.style = 10
    chart.title = f"Fees (CHF) by {label}"
    chart.legend = None
    chart.height, chart.width = 9, 18
    data = ChartRef(ws, min_col=fees_col, min_row=rows[0] - 1, max_row=rows[1])
    cats = ChartRef(ws, min_col=1, min_row=rows[0], max_row=rows[1])
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.series[0].graphicalProperties.solidFill = excel.TIGER_ORANGE
    if label != "Month":
        chart.y_axis.scaling.orientation = "minMax"
        chart.x_axis.scaling.orientation = "maxMin"     # largest at the top
    ws.add_chart(chart, anchor)
