"""Inbox -> clean -> register -> reports, in one call.

Workspace layout (created on first run, never committed to git):

    workspace/
      inbox/                 drop trade files here (xlsx, csv, html)
      inbox/processed/       files are moved here once read
      register.xlsx          the master transaction register
      output/
        exceptions_*.xlsx    rows refused, with the reason; fix and drop back
        import_log_*.xlsx    what was accepted, corrected and refused
        reoffer_report.xlsx  dashboards, rebuilt on every run
"""

import datetime as dt
import shutil
from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd
from openpyxl import Workbook

from . import excel
from .clean import blank, clean_frame, dedupe_key
from .ingest import SUPPORTED, read_file
from .reference import Reference
from .schema import HEADERS, KEYS

REGISTER_COLUMNS = ["trade_id"] + KEYS + ["source_file", "imported_at", "warnings"]
FORMATS = {"Trade date": "date", "Size": "amount", "Fees (%)": "pct",
           "Fees (Amount)": "amount2", "Imported at": "date"}


@dataclass
class Workspace:
    root: Path

    def __post_init__(self):
        self.root = Path(self.root)
        self.inbox = self.root / "inbox"
        self.processed = self.inbox / "processed"
        self.output = self.root / "output"
        self.register_path = self.root / "register.xlsx"
        for p in (self.inbox, self.processed, self.output):
            p.mkdir(parents=True, exist_ok=True)


@dataclass
class RunResult:
    files: list = field(default_factory=list)
    accepted: int = 0
    corrected: int = 0
    rejected: int = 0
    duplicates: int = 0
    ignored: int = 0
    file_errors: list = field(default_factory=list)
    exceptions_path: Path = None
    log_path: Path = None
    report_path: Path = None


# --------------------------------------------------------------------------
# Register I/O
# --------------------------------------------------------------------------

def load_register(path) -> pd.DataFrame:
    if not Path(path).exists():
        return pd.DataFrame(columns=REGISTER_COLUMNS)
    df = pd.read_excel(path, sheet_name="Register", dtype={
        "ISIN": str, "nclidos": str, "Trade ID": str})
    inv = {"Trade ID": "trade_id", "Source file": "source_file",
           "Imported at": "imported_at", "Warnings": "warnings",
           **{h: k for k, h in HEADERS.items()}}
    df = df.rename(columns=inv)
    df["trade_date"] = pd.to_datetime(df["trade_date"]).dt.date
    for c in ("nclidos", "isin", "comments", "warnings"):
        df[c] = df[c].fillna("").astype(str)
    return df[REGISTER_COLUMNS]


def save_register(df, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Register"
    out = df[REGISTER_COLUMNS].rename(columns={
        "trade_id": "Trade ID", "source_file": "Source file",
        "imported_at": "Imported at", "warnings": "Warnings", **HEADERS})
    excel.write_table(ws, out, formats=FORMATS, name="Register",
                      widths={"Warnings": 50, "Comments": 40, "Source file": 30})
    ws.freeze_panes = "C2"
    tmp = Path(path).with_suffix(".tmp.xlsx")
    wb.save(tmp)
    tmp.replace(path)                     # never leave a half-written register


# --------------------------------------------------------------------------
# Run
# --------------------------------------------------------------------------

def run(workspace, reference_dir=None, files=None, today=None, build_report=True):
    """Process every file in the inbox (or `files`) and refresh the reports."""
    ws = Workspace(workspace)
    ref = Reference.load(reference_dir) if reference_dir else Reference.load()
    today = today or dt.date.today()
    stamp = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
    res = RunResult()

    paths = [Path(f) for f in files] if files else sorted(
        p for p in ws.inbox.iterdir() if p.is_file() and p.suffix.lower() in SUPPORTED
        and not p.name.startswith("~$"))

    raws = []
    for p in paths:
        try:
            raws.append(read_file(p))
            res.files.append(p)
        except Exception as e:                       # one bad file must not stop the run
            res.file_errors.append(f"{p.name}: {e}")
    register = load_register(ws.register_path)

    if raws:
        raw = pd.concat(raws, ignore_index=True)
        # Rows with neither ISIN nor size are notes / reference blocks, not trades.
        is_trade = ~(raw["isin"].map(blank) & raw["size"].map(blank))
        res.ignored = int((~is_trade).sum())
        accepted, rejected = clean_frame(raw[is_trade], ref, today)

        # Duplicates: inside this batch and against the register.
        if len(accepted):
            keys = dedupe_key(accepted)
            known = dict(zip(dedupe_key(register), register["trade_id"])) if len(register) else {}
            dup_reg = keys.map(lambda k: k in known)
            dup_batch = keys.duplicated() & ~dup_reg
            dup = dup_reg | dup_batch
            if dup.any():
                d = accepted[dup].copy()
                d["errors"] = [f"duplicate of register trade {known[k]}" if k in known
                               else "duplicate row in this import"
                               for k in keys[dup]]
                rejected = pd.concat([rejected, d.drop(columns="warnings")], ignore_index=True)
                res.duplicates = int(dup.sum())
                accepted = accepted[~dup]

        res.accepted, res.rejected = len(accepted), len(rejected)
        res.corrected = int((accepted["warnings"] != "").sum())

        if len(accepted):
            start = _next_id(register)
            accepted = accepted.sort_values("trade_date", kind="stable").copy()
            accepted["trade_id"] = [f"R{n:06d}" for n in range(start, start + len(accepted))]
            accepted["imported_at"] = today
            register = pd.concat([register, accepted[REGISTER_COLUMNS]], ignore_index=True)
            register = register.sort_values(["trade_date", "trade_id"], kind="stable")
            save_register(register, ws.register_path)

        if len(rejected):
            res.exceptions_path = ws.output / f"exceptions_{stamp}.xlsx"
            write_exceptions(rejected, res.exceptions_path)
        res.log_path = ws.output / f"import_log_{stamp}.xlsx"
        write_log(accepted, rejected, res, res.log_path)

        if not files:
            dest = ws.processed / today.isoformat()
            dest.mkdir(parents=True, exist_ok=True)
            for p in res.files:
                shutil.move(str(p), dest / p.name)

    if build_report and len(register):
        from .report import build_report as _build
        res.report_path = ws.output / "reoffer_report.xlsx"
        _build(register, ref, res.report_path)
    return res


def _next_id(register):
    if not len(register):
        return 1
    nums = register["trade_id"].astype(str).str.extract(r"(\d+)")[0].dropna().astype(int)
    return int(nums.max()) + 1 if len(nums) else 1


def write_exceptions(rejected, path):
    """Same columns as the input template, plus the reason, so the file can be
    corrected and dropped straight back into the inbox."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Exceptions"
    out = rejected[KEYS + ["errors", "source_file", "source_row"]].rename(columns={
        **HEADERS, "errors": "Errors (fix, then re-drop this file)",
        "source_file": "From file", "source_row": "From row"})
    excel.write_table(ws, out, name="Exceptions",
                      widths={"Errors (fix, then re-drop this file)": 70})
    err_col = len(KEYS) + 1
    for r in range(2, len(out) + 2):
        ws.cell(r, err_col).fill = excel.PatternFill("solid", fgColor=excel.ERROR_FILL)
    ws.freeze_panes = "A2"
    wb.save(path)


def write_log(accepted, rejected, res, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    excel.title(ws, "Reoffer import log", dt.datetime.now().strftime("%d.%m.%Y %H:%M"))
    rows = [("Files read", len(res.files)), ("Rows accepted", res.accepted),
            ("  of which auto-corrected", res.corrected),
            ("Rows refused", res.rejected), ("  of which duplicates", res.duplicates),
            ("Non-trade rows skipped", res.ignored)]
    rows += [("File not read", e) for e in res.file_errors]
    excel.write_table(ws, pd.DataFrame(rows, columns=["Item", "Value"]), start_row=4,
                      widths={"Item": 30, "Value": 60})
    if len(accepted):
        w = wb.create_sheet("Corrected")
        c = accepted[accepted["warnings"] != ""]
        out = c[["source_file", "source_row", "isin", "nclidos", "warnings"]].rename(
            columns={"source_file": "File", "source_row": "Row", "isin": "ISIN",
                     "warnings": "What was changed"})
        excel.write_table(w, out, widths={"What was changed": 80})
    if len(rejected):
        w = wb.create_sheet("Refused")
        out = rejected[["source_file", "source_row", "isin", "nclidos", "errors"]].rename(
            columns={"source_file": "File", "source_row": "Row", "isin": "ISIN",
                     "errors": "Why"})
        excel.write_table(w, out, widths={"Why": 80})
    wb.save(path)
