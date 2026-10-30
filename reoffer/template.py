"""Generate the Excel input template with dropdowns built from the reference
lists, so most typos cannot be typed in the first place."""

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from . import excel
from .reference import Reference
from .schema import COLUMNS, HEADERS, KEYS, LIST_FIELDS

ROWS = 2000
HELP = {
    "trade_date": "Date of the trade, DD.MM.YYYY",
    "isin": "12 characters, e.g. CH1234567890. Check digit is verified.",
    "nclidos": "Client / account ID, kept as text",
    "size": "Notional in the trade currency",
    "fee_pct": "Fee rate, e.g. 1.84%",
    "fee_amount": "Leave the formula: Size x Fees (%). Override only with a comment.",
    "comments": "Required when the fee amount is overridden (invoice adjustment)",
}


def build_template(path, ref=None):
    ref = ref or Reference.load()
    wb = Workbook()
    ws = wb.active
    ws.title = "Trades"
    lists = wb.create_sheet("Lists")

    for j, (key, header, required) in enumerate(COLUMNS, 1):
        c = ws.cell(1, j, header)
        c.font, c.fill = excel.HEADER_FONT, excel.HEADER_FILL
        if not required:
            c.fill = excel.PatternFill("solid", fgColor=excel.MUTED_INK)
        if key in HELP:
            c.comment = Comment(HELP[key], "Reoffer")
        ws.column_dimensions[get_column_letter(j)].width = 22 if key != "comments" else 45

    col = {k: get_column_letter(KEYS.index(k) + 1) for k in KEYS}
    rng = lambda k: f"{col[k]}2:{col[k]}{ROWS + 1}"

    # Dropdowns, one column per list on the hidden Lists sheet.
    for j, key in enumerate(LIST_FIELDS, 1):
        values = ref.lists.get(key, [])
        lists.cell(1, j, HEADERS[key])
        for i, v in enumerate(values, 2):
            lists.cell(i, j, v)
        if values:
            L = get_column_letter(j)
            dv = DataValidation(type="list", formula1=f"=Lists!${L}$2:${L}${len(values) + 1}",
                                allow_blank=True, showErrorMessage=True,
                                errorTitle="Not in the list",
                                error=f"Pick a {HEADERS[key]} from the list. New values "
                                      "go in reference/lists.csv.")
            ws.add_data_validation(dv)
            dv.add(rng(key))
    lists.sheet_state = "hidden"

    rules = [
        ("trade_date", DataValidation(type="date", operator="greaterThan", formula1="36526",
                                      error="Enter a date (DD.MM.YYYY)")),
        ("isin", DataValidation(type="textLength", operator="equal", formula1="12",
                                error="An ISIN has exactly 12 characters")),
        ("size", DataValidation(type="decimal", operator="greaterThan", formula1="0",
                                error="Size must be a positive number")),
        ("fee_pct", DataValidation(type="decimal", operator="between", formula1="0",
                                   formula2="0.06", error="Fee rate between 0% and 6%")),
    ]
    for key, dv in rules:
        dv.showErrorMessage = True
        ws.add_data_validation(dv)
        dv.add(rng(key))

    for r in range(2, ROWS + 2):
        ws[f"{col['trade_date']}{r}"].number_format = excel.FMT["date"]
        ws[f"{col['nclidos']}{r}"].number_format = excel.FMT["text"]
        ws[f"{col['size']}{r}"].number_format = excel.FMT["amount"]
        ws[f"{col['fee_pct']}{r}"].number_format = excel.FMT["pct"]
        f = ws[f"{col['fee_amount']}{r}"]
        f.value = (f'=IF(AND({col["size"]}{r}<>"",{col["fee_pct"]}{r}<>""),'
                   f'ROUND({col["size"]}{r}*{col["fee_pct"]}{r},2),"")')
        f.number_format = excel.FMT["amount2"]
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(KEYS))}1"
    wb.save(path)
    return path
