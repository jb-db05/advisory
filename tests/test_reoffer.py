import datetime as dt

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from reoffer.clean import isin_valid, parse_fee_rate, parse_number
from reoffer.pipeline import load_register, run
from reoffer.reference import Reference
from reoffer.schema import HEADERS, KEYS
from reoffer.template import build_template

TODAY = dt.date(2026, 10, 30)


def make_isin(body):
    """Append the check digit to an 11-character ISIN body."""
    for d in "0123456789":
        if isin_valid(body + d):
            return body + d


ISIN_A = make_isin("CH131234567")
ISIN_B = make_isin("XS281234567")


@pytest.fixture
def ref(tmp_path):
    d = tmp_path / "ref"
    d.mkdir()
    (d / "lists.csv").write_text(
        "field,value\nissuer,UBS\nissuer,BNP Paribas\nissuer,Goldman Sachs\n"
        "payoff,Barrier Reverse Convertible\npayoff,Autocallable\n"
        "currency,CHF\ncurrency,USD\ncurrency,EUR\napl,Yes\napl,No\n"
        "rm,Anna Keller\nrm,Marc Dubois\nregion,Switzerland\nregion,Middle East\n"
        "management_style,Advisory\nmanagement_style,Discretionary\n", encoding="utf-8")
    (d / "aliases.csv").write_text(
        "field,alias,canonical\nissuer,GS,Goldman Sachs\npayoff,BRC,Barrier Reverse Convertible\n",
        encoding="utf-8")
    (d / "fx_rates.csv").write_text(
        "date,currency,chf_per_unit\n2026-01-01,USD,0.90\n2026-01-01,EUR,0.95\n",
        encoding="utf-8")
    return d


def trade(**kw):
    base = dict(trade_date=dt.datetime(2026, 3, 2), isin=ISIN_A, issuer="UBS",
                payoff="Autocallable", nclidos="100234", size=130000, currency="USD",
                fee_pct=0.0184, fee_amount=2392, apl="Yes", rm="Anna Keller",
                region="Switzerland", management_style="Advisory", comments=None)
    base.update(kw)
    return [base[k] for k in KEYS]


def legacy_workbook(path, rows):
    """Header in row 1, reference block in rows 2-22, trades from row 23."""
    wb = Workbook()
    ws = wb.active
    ws.append([HEADERS[k] for k in KEYS])
    for i in range(21):
        ws.append([None, None, None, None, None, None, None, None, None,
                   "Yes" if i % 2 else "No", None, "Switzerland"])
    for r in rows:
        ws.append(r)
    wb.save(path)


# ---------------------------------------------------------------- parsers

@pytest.mark.parametrize("raw,expected", [
    ("130'000", 130000), ("130,000", 130000), ("1.234,56", 1234.56),
    ("1,234.56", 1234.56), ("USD 250000", 250000), ("1.5m", 1_500_000),
    ("500k", 500000), (130000, 130000),
])
def test_parse_number(raw, expected):
    assert parse_number(raw) == pytest.approx(expected)


@pytest.mark.parametrize("raw,expected", [
    ("1.84%", 0.0184), ("1,84", 0.0184), (1.84, 0.0184), (0.0184, 0.0184),
])
def test_parse_fee_rate(raw, expected):
    assert parse_fee_rate(raw)[0] == pytest.approx(expected)


def test_fee_rate_disambiguated_by_amount():
    # 0.04 alone reads as 4%; with a 40 fee on 100k it must be 0.04%.
    assert parse_fee_rate(0.04, 100000, 40)[0] == pytest.approx(0.0004)


def test_isin_check_digit():
    assert isin_valid("US0378331005")          # Apple
    assert not isin_valid("US0378331006")
    assert not isin_valid("US037833100")


def test_resolve(ref):
    r = Reference.load(ref)
    assert r.resolve("issuer", "ubs") == ("UBS", "")
    assert r.resolve("issuer", "GS") == ("Goldman Sachs", "")
    assert r.resolve("rm", "Ana Keller") == ("Anna Keller", "corrected")
    canon, msg = r.resolve("issuer", "Barclays")
    assert canon is None and "not a known issuer" in msg


# ---------------------------------------------------------------- end to end

def test_run_legacy_csv_html(tmp_path, ref):
    ws = tmp_path / "ws"
    inbox = ws / "inbox"
    inbox.mkdir(parents=True)
    legacy_workbook(inbox / "Reoffer 2026.xlsx", [
        trade(),                                             # clean
        trade(isin=ISIN_B, issuer="gs", payoff="BRC", rm="Ana Keller",
              size="250'000", currency="eur", fee_pct="1,5%", fee_amount=None),
        trade(nclidos="100999", fee_amount=3000),            # fee mismatch, no comment
        trade(nclidos="100998", fee_amount=3000, comments="invoice adj. -608"),
        trade(isin="CH0000000000"),                          # bad ISIN
        trade(issuer="Barclays"),                            # unknown issuer
        trade(),                                             # duplicate of row 1
    ])
    (inbox / "desk.csv").write_text(
        "Reoffer export\n\n"
        "Trade date;ISIN;Issuer;Pay-off;nclidos;Notional;Ccy;Fees (%);RM;Region;"
        "Management Style;APL\n"
        f"15.04.2026;{ISIN_B};BNP Paribas;Autocallable;200111;1'000'000;CHF;0,75;"
        "Marc Dubois;Middle East;Discretionary;No\n", encoding="cp1252")
    (inbox / "confirm.html").write_text(
        "<html><body><h1>Confirmation</h1><table>"
        "<tr><th>Trade date</th><th>ISIN</th><th>Issuer</th><th>Payoff</th><th>Client ID</th>"
        "<th>Size</th><th>Currency</th><th>Fees</th><th>APL</th><th>RM</th>"
        "<th>Region</th><th>Mandate</th></tr>"
        f"<tr><td>2026-05-20</td><td>{ISIN_A}</td><td>UBS</td><td>Autocallable</td>"
        "<td>300500</td><td>500,000</td><td>USD</td><td>1.2%</td><td>Yes</td>"
        "<td>Marc Dubois</td><td>Switzerland</td><td>Advisory</td></tr>"
        "</table></body></html>", encoding="utf-8")

    res = run(ws, ref, today=TODAY)
    assert res.file_errors == []
    assert res.ignored == 21                     # reference block skipped silently
    assert res.accepted == 5
    assert res.rejected == 4 and res.duplicates == 1

    reg = load_register(ws / "register.xlsx")
    assert reg["trade_id"].is_unique
    brc = reg[reg["isin"] == ISIN_B].iloc[0]
    assert brc["issuer"] == "Goldman Sachs" and brc["rm"] == "Anna Keller"
    assert brc["currency"] == "EUR" and brc["fee_amount"] == pytest.approx(3750)
    assert "corrected" in brc["warnings"]
    assert set(reg["nclidos"]) >= {"100234", "200111", "300500"}

    exc = pd.read_excel(res.exceptions_path)
    errs = " | ".join(exc.iloc[:, len(KEYS)].astype(str))
    assert "fix it or explain in Comments" in errs
    assert "invalid" in errs and "not a known issuer" in errs and "duplicate" in errs

    # Inbox emptied, report built with CHF figures.
    assert not any(p.is_file() for p in inbox.iterdir())
    wb = load_workbook(res.report_path)
    assert {"Summary", "By Month", "By RM", "RM x Month", "Data"} <= set(wb.sheetnames)
    data = pd.read_excel(res.report_path, sheet_name="Data")
    usd = data[data["nclidos"].astype(str) == "100234"].iloc[0]
    assert usd["Fees (CHF)"] == pytest.approx(2392 * 0.90)

    # Re-importing the same trades is refused as duplicates, register unchanged.
    legacy_workbook(inbox / "again.xlsx", [trade()])
    res2 = run(ws, ref, today=TODAY)
    assert res2.accepted == 0 and res2.duplicates == 1
    assert len(load_register(ws / "register.xlsx")) == 5


def test_exceptions_file_can_be_redropped(tmp_path, ref):
    ws = tmp_path / "ws"
    (ws / "inbox").mkdir(parents=True)
    legacy_workbook(ws / "inbox" / "a.xlsx", [trade(issuer="Barclays")])
    res = run(ws, ref, today=TODAY)
    assert res.rejected == 1
    exc = pd.read_excel(res.exceptions_path)
    exc.loc[0, "Issuer"] = "UBS"                      # the user's fix
    exc.to_excel(ws / "inbox" / "fixed.xlsx", index=False)
    res2 = run(ws, ref, today=TODAY)
    assert res2.accepted == 1 and res2.rejected == 0


def test_missing_fx_is_reported(tmp_path, ref):
    ws = tmp_path / "ws"
    (ws / "inbox").mkdir(parents=True)
    (ref / "lists.csv").write_text((ref / "lists.csv").read_text() + "currency,GBP\n")
    legacy_workbook(ws / "inbox" / "a.xlsx", [trade(currency="GBP")])
    res = run(ws, ref, today=TODAY)
    assert "Missing FX" in load_workbook(res.report_path).sheetnames


def test_template(tmp_path, ref):
    p = build_template(tmp_path / "t.xlsx", Reference.load(ref))
    wb = load_workbook(p)
    ws = wb["Trades"]
    assert [c.value for c in ws[1]] == [HEADERS[k] for k in KEYS]
    assert ws["I2"].value.startswith("=IF(")
    assert wb["Lists"].sheet_state == "hidden"
    assert len(ws.data_validations.dataValidation) >= 7
